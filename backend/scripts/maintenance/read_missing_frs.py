"""
Lit les 6 FRs docx non structurées et affiche leur contenu textuel.
"""
import docx
from pathlib import Path

FR_DIR = Path(__file__).parent / "FR"

FILES = [
    "FR 077  CARACTERES ERRONES EN BASE BRASIL champ t_ports.t_remarks.docx",
    "FR 114 - PORT RESEAU SANS SERVICE NI EXTREMITE.docx",
    "FR 165 Carte impossible \u00e0 supprimer.docx",
    "FR 176 Modification de l'\u00e9tat d'un TP.docx",
    "FR 182 G8 IHM BLOQUEE SUR DS PARAM.docx",
    "FR 203 - FAUSSE MANIP BC ( Perte Des Donn\u00e9es ).docx",
]

for fname in FILES:
    path = FR_DIR / fname
    if not path.exists():
        print(f"NOT FOUND: {fname}")
        continue
    print(f"\n{'='*70}")
    print(f"  {fname}")
    print(f"{'='*70}")
    try:
        doc = docx.Document(str(path))
        for para in doc.paragraphs:
            txt = para.text.strip()
            if txt:
                print(txt)
        # Tables
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    print(" | ".join(cells))
    except Exception as e:
        print(f"  ERROR: {e}")
