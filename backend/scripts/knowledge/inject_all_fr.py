"""
inject_all_fr.py
Indexe TOUTES les FRs disponibles dans backend/FR/*.docx dans Qdrant brasil_procedures.

- Extraction automatique du numéro FR depuis le nom de fichier
- Parsing intelligent du contenu docx (symptômes, étapes, SQL, tables, causes)
- Embedding réel (paraphrase-multilingual-MiniLM-L12-v2)
- UUID déterministe par proc_id → idempotent (relancer = pas de doublons)
- Mode --force pour ré-indexer même les points existants

Usage :
    python inject_all_fr.py           # indexe les manquants seulement
    python inject_all_fr.py --force   # ré-indexe tout
    python inject_all_fr.py --list    # liste les FRs détectées sans indexer
"""
import sys, os, glob, uuid, re
from datetime import datetime

try:
    import docx
except ImportError:
    print("pip install python-docx"); sys.exit(1)
try:
    from qdrant_client import QdrantClient
    from qdrant_client.models import PointStruct
except ImportError:
    print("pip install qdrant-client"); sys.exit(1)

# ─────────────────────────────────────────────────────────────────────────────
COLLECTION = "brasil_procedures"
FR_DIR = os.path.join(os.path.dirname(__file__), "FR")
EMBED_DIM = 384

# ─────────────────────────────────────────────────────────────────────────────
# Modèle embedding (chargé une seule fois)
# ─────────────────────────────────────────────────────────────────────────────
_model = None

def get_model():
    global _model
    if _model is not None:
        return _model
    try:
        from sentence_transformers import SentenceTransformer
        hf_base = os.path.expanduser(
            "~/.cache/huggingface/hub/"
            "models--sentence-transformers--paraphrase-multilingual-MiniLM-L12-v2"
        )
        snaps = glob.glob(os.path.join(hf_base, "snapshots", "*"))
        path = snaps[0] if snaps else "paraphrase-multilingual-MiniLM-L12-v2"
        _model = SentenceTransformer(path, device="cpu")
        print(f"[Embed] Modèle chargé depuis: {path}")
    except Exception as e:
        print(f"[Embed] WARN modèle indisponible ({e}) — embeddings zéro")
        _model = None
    return _model

def embed(text: str) -> list:
    m = get_model()
    if m is None:
        return [0.0] * EMBED_DIM
    return m.encode(text, show_progress_bar=False).tolist()


# ─────────────────────────────────────────────────────────────────────────────
# Extraction du numéro FR depuis le nom de fichier
# ─────────────────────────────────────────────────────────────────────────────
_FR_NUM_RE = re.compile(r"FR[\s_-]?(\d+[A-Za-z]?(?:\d+)?)", re.IGNORECASE)

def extract_fr_number(filename: str) -> str | None:
    """Retourne '001', '136b', '1583B', etc. depuis le nom de fichier."""
    m = _FR_NUM_RE.search(os.path.basename(filename))
    return m.group(1) if m else None


# ─────────────────────────────────────────────────────────────────────────────
# Extraction du contenu docx
# ─────────────────────────────────────────────────────────────────────────────
_SQL_RE = re.compile(
    r"\b(SELECT|INSERT|UPDATE|DELETE|FROM|WHERE|JOIN|GROUP BY|ORDER BY"
    r"|CREATE|ALTER|DROP|TRUNCATE|WITH|HAVING|UNION|MERGE)\b",
    re.IGNORECASE,
)
_TABLE_RE = re.compile(r"\b(t_[a-z_]+|lst_[a-z_]+|v_[a-z_]+)\b", re.IGNORECASE)
_EXCEPTION_RE = re.compile(r"\b[A-Z][a-z]+(?:[A-Z][a-z]+)+Exception\b")

def parse_docx(path: str) -> dict:
    """
    Extrait depuis le docx:
    - full_text, paragraphs
    - title (premier paragraphe non vide)
    - symptoms, diagnostic_steps, resolution_steps, root_causes
    - sql_queries, tables_involved, exceptions_referenced
    """
    doc = docx.Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Texte des tableaux
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                tables_text.append(" | ".join(cells))

    all_lines = paragraphs + tables_text
    full_text = "\n".join(all_lines)

    # ── Titre ────────────────────────────────────────────────────────────────
    title = paragraphs[0] if paragraphs else os.path.basename(path).replace(".docx", "")

    # ── Extraction des sections par mots-clés ─────────────────────────────────
    # On classe chaque ligne dans une section selon les titres de section courants
    _SECTION_KEYS = {
        "symptom":   re.compile(r"symptôme|symptome|contexte|problème|probl[eè]me|incident|description|objet|sujet", re.I),
        "cause":     re.compile(r"cause|origine|raison|root.cause|analyse", re.I),
        "diag":      re.compile(r"diagnostic|vérif|verif|contrôle|controle|étape.*diagnos|requête.*diagnos", re.I),
        "resolution":re.compile(r"résolution|resolution|solution|procédure|procedure|action|correction|traitement|étape", re.I),
    }

    sections: dict[str, list] = {"symptom": [], "cause": [], "diag": [], "resolution": [], "other": []}
    current_section = "other"

    for line in all_lines:
        # Détecter un titre de section (ligne courte, souvent en majuscules ou avec :)
        is_heading = (
            len(line) < 80
            and (line.isupper() or line.endswith(":") or any(k.search(line) for k in _SECTION_KEYS.values()))
        )
        if is_heading:
            matched = False
            for sec, pat in _SECTION_KEYS.items():
                if pat.search(line):
                    current_section = sec
                    matched = True
                    break
            if not matched and is_heading and len(line) < 60:
                # Titre inconnu — ne pas changer de section
                pass
        else:
            sections[current_section].append(line)

    # Nettoyer : pas de lignes trop courtes (< 10 chars) ni doublons
    def clean(lst): 
        seen = set()
        out = []
        for l in lst:
            s = l.strip()
            if len(s) >= 10 and s not in seen:
                seen.add(s)
                out.append(s)
        return out[:12]  # max 12 items par section

    symptoms         = clean(sections["symptom"]) or clean(sections["other"][:3])
    root_causes      = clean(sections["cause"])
    diagnostic_steps = clean(sections["diag"])
    resolution_steps = clean(sections["resolution"])

    # Si resolution_steps vide, prendre "other" sauf les 3 premières lignes (déjà dans symptoms)
    if not resolution_steps:
        resolution_steps = clean(sections["other"][3:])

    # ── SQL queries ───────────────────────────────────────────────────────────
    sql_queries = []
    _buf = []
    for line in all_lines:
        if _SQL_RE.search(line) and len(line) > 15:
            _buf.append(line.strip())
        else:
            if _buf:
                sql_queries.append(" ".join(_buf))
                _buf = []
    if _buf:
        sql_queries.append(" ".join(_buf))
    # Dédup + max 10
    seen_sql = set()
    sql_out = []
    for q in sql_queries:
        if q not in seen_sql:
            seen_sql.add(q)
            sql_out.append(q)
    sql_queries = sql_out[:10]

    # ── Tables BRASIL ─────────────────────────────────────────────────────────
    tables_found = set()
    for line in all_lines:
        for m in _TABLE_RE.finditer(line):
            tables_found.add(m.group(0).lower())
    tables_involved = sorted(tables_found)[:20]

    # ── Exceptions Java ───────────────────────────────────────────────────────
    exceptions_found = set()
    for line in all_lines:
        for m in _EXCEPTION_RE.finditer(line):
            exceptions_found.add(m.group(0))
    exceptions_referenced = sorted(exceptions_found)

    return {
        "title": title,
        "full_text": full_text,
        "symptoms": symptoms,
        "root_causes": root_causes,
        "diagnostic_steps": diagnostic_steps,
        "resolution_steps": resolution_steps,
        "sql_queries": sql_queries,
        "tables_involved": tables_involved,
        "exceptions_referenced": exceptions_referenced,
    }


# ─────────────────────────────────────────────────────────────────────────────
# Construire un texte d'embedding représentatif
# ─────────────────────────────────────────────────────────────────────────────
def build_embed_text(fr_label: str, parsed: dict) -> str:
    parts = [fr_label, parsed["title"]]
    if parsed["symptoms"]:
        parts.append("Symptômes: " + ". ".join(parsed["symptoms"][:3]))
    if parsed["root_causes"]:
        parts.append("Causes: " + ". ".join(parsed["root_causes"][:2]))
    if parsed["tables_involved"]:
        parts.append("Tables: " + " ".join(parsed["tables_involved"][:8]))
    if parsed["exceptions_referenced"]:
        parts.append(" ".join(parsed["exceptions_referenced"][:4]))
    return " ".join(parts)[:512]


# ─────────────────────────────────────────────────────────────────────────────
# Indexation d'un fichier FR
# ─────────────────────────────────────────────────────────────────────────────
def index_fr(client: QdrantClient, path: str, force: bool = False) -> str:
    """
    Indexe un fichier FR dans Qdrant.
    Retourne "indexed", "skipped", "error"
    """
    filename = os.path.basename(path)
    fr_raw = extract_fr_number(filename)
    if not fr_raw:
        return "error"

    fr_label   = f"FR {fr_raw}"
    proc_id    = f"FR-{fr_raw.upper()}-AUTO"
    point_uuid = str(uuid.uuid5(uuid.NAMESPACE_DNS, proc_id))

    # Vérifier doublons sauf si --force
    if not force:
        try:
            existing = client.retrieve(
                collection_name=COLLECTION,
                ids=[point_uuid],
                with_payload=False,
            )
            if existing:
                print(f"  [SKIP] {fr_label} déjà indexée (UUID={point_uuid})")
                return "skipped"
        except Exception:
            pass

    # Parser le docx
    try:
        parsed = parse_docx(path)
    except Exception as e:
        print(f"  [ERROR] Lecture {filename}: {e}")
        return "error"

    # Construire le payload
    payload = {
        "procedure_id":        proc_id,
        "fr_number":           fr_label,
        "fr_aliases":          [f"FR{fr_raw}", f"FR-{fr_raw}", fr_label],
        "title":               parsed["title"],
        "title_normalized":    parsed["title"].lower(),
        "application":         "BRASIL",
        "applications_involved": ["BRASIL"],
        "symptoms":            parsed["symptoms"],
        "diagnostic_steps":    parsed["diagnostic_steps"],
        "root_causes":         parsed["root_causes"],
        "resolution_steps":    parsed["resolution_steps"],
        "sql_queries":         parsed["sql_queries"],
        "tables_involved":     parsed["tables_involved"],
        "exceptions_referenced": parsed["exceptions_referenced"],
        "source_types":        ["fr_document"],
        "trust_score":         0.85,
        "ticket_count":        0,
        "cluster_id":          f"FR-{fr_raw.upper()}",
        "full_text_excerpt":   parsed["full_text"][:2000],
        "indexed_at":          datetime.now().isoformat(),
        "last_updated":        datetime.now().isoformat(),
    }

    # Embedding
    embed_text = build_embed_text(fr_label, parsed)
    vector = embed(embed_text)

    # Upsert
    try:
        result = client.upsert(
            collection_name=COLLECTION,
            points=[PointStruct(id=point_uuid, vector=vector, payload=payload)],
        )
        print(f"  [OK] {fr_label} — {parsed['title'][:60]}")
        print(f"       tables={len(parsed['tables_involved'])}  sql={len(parsed['sql_queries'])}  steps={len(parsed['resolution_steps'])}")
        return "indexed"
    except Exception as e:
        print(f"  [ERROR] Upsert {fr_label}: {e}")
        return "error"


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────
def main():
    force = "--force" in sys.argv
    list_only = "--list" in sys.argv

    print("=" * 65)
    print("inject_all_fr.py — Indexation de toutes les FRs dans Qdrant")
    print(f"Mode: {'FORCE (ré-index tout)' if force else 'INCREMENTAL (saute les existantes)'}")
    print("=" * 65)

    # Lister tous les docx
    all_files = sorted(glob.glob(os.path.join(FR_DIR, "*.docx")))
    print(f"\n{len(all_files)} fichiers .docx trouvés dans {FR_DIR}\n")

    if list_only:
        for f in all_files:
            num = extract_fr_number(f)
            print(f"  {os.path.basename(f)}  →  FR {num}")
        return

    # Connexion Qdrant
    client = QdrantClient(host="localhost", port=6333)
    try:
        info = client.get_collection(COLLECTION)
        print(f"Collection {COLLECTION}: {info.points_count} points existants\n")
    except Exception as e:
        print(f"[ERROR] Collection {COLLECTION} inaccessible: {e}")
        sys.exit(1)

    # Charger le modèle une fois
    get_model()
    print()

    # Indexer
    counts = {"indexed": 0, "skipped": 0, "error": 0}
    for i, path in enumerate(all_files, 1):
        print(f"[{i:02d}/{len(all_files)}] {os.path.basename(path)}")
        status = index_fr(client, path, force=force)
        counts[status] += 1

    # Résumé
    info = client.get_collection(COLLECTION)
    print(f"\n{'='*65}")
    print(f"Résultat final:")
    print(f"  Indexées  : {counts['indexed']}")
    print(f"  Ignorées  : {counts['skipped']} (déjà présentes)")
    print(f"  Erreurs   : {counts['error']}")
    print(f"  Total collection brasil_procedures: {info.points_count} points")
    print("=" * 65)


if __name__ == "__main__":
    main()
