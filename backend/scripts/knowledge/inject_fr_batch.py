"""
inject_fr_batch.py
Indexe en batch plusieurs FRs dans Qdrant brasil_procedures.
FRs ciblées : FR 188, FR 190, FR 191 (+ extensible).

Usage :
    python inject_fr_batch.py
"""
import sys, os, glob, uuid, re
from datetime import datetime

try:
    import docx
except ImportError:
    print("pip install python-docx"); sys.exit(1)
try:
    from qdrant_client import QdrantClient
    from qdrant_client.models import PointStruct
except ImportError:
    print("pip install qdrant-client"); sys.exit(1)

FR_DIR = os.path.join(os.path.dirname(__file__), "FR")

# ── Modèle d'embedding ────────────────────────────────────────────────────────
_embedding_model = None

def get_embedding_model():
    global _embedding_model
    if _embedding_model is not None:
        return _embedding_model
    try:
        from sentence_transformers import SentenceTransformer
        import glob as _g
        hf_hub_base = os.path.expanduser(
            '~/.cache/huggingface/hub/models--sentence-transformers--paraphrase-multilingual-MiniLM-L12-v2'
        )
        snapshots = _g.glob(os.path.join(hf_hub_base, 'snapshots', '*'))
        if snapshots:
            _embedding_model = SentenceTransformer(snapshots[0], device='cpu')
        else:
            _embedding_model = SentenceTransformer('paraphrase-multilingual-MiniLM-L12-v2', device='cpu')
        print(f"[OK] Modèle embedding chargé")
    except Exception as e:
        print(f"[WARN] SentenceTransformer indisponible ({e}), embeddings zéro")
        _embedding_model = None
    return _embedding_model


def compute_embedding(text: str) -> list:
    model = get_embedding_model()
    if model is None:
        return [0.0] * 384
    return model.encode(text).tolist()


def extract_docx(path: str) -> tuple[str, list, list]:
    """Retourne (full_text, paragraphs, tables_text)"""
    doc = docx.Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                tables_text.append(" | ".join(cells))
    full_text = "\n".join(paragraphs + tables_text)
    return full_text, paragraphs, tables_text


# ── Définitions des FRs à indexer ─────────────────────────────────────────────

FR_DEFINITIONS = [

    # ── FR 190 : Suppression VLAN IMPOSSIBLE ──────────────────────────────────
    {
        "fr_number": "FR 190",
        "fr_aliases": ["FR190", "FR-190", "FR 190"],
        "glob_pattern": "*190*VLAN*",
        "proc_id": "FR-190-SUPPRESSION-VLAN",
        "title": "FR 190 - Suppression VLAN IMPOSSIBLE",
        "title_normalized": "suppression vlan impossible",
        "incident_type": "network_resource.vlan_deletion_impossible",
        "application": "BRASIL",
        "applications_involved": ["BRASIL"],
        "estimated_duration": "15 à 30 minutes",
        "sensitivity": "HAUTE — vérifier les associations VLAN avant suppression",
        "symptoms": [
            "Suppression d'un VLAN via IHM BRASIL impossible",
            "Message d'erreur lors de la tentative de suppression d'un VLAN",
            "VLAN encore associé à des équipements ou des ressources logiques",
            "lst_vlan_usage retourne des lignes pour ce VLAN",
        ],
        "diagnostic_steps": [
            "Identifier l'identifiant du VLAN dans la base : SELECT vlan_id, vlan_name FROM t_vlans WHERE vlan_name = 'NOM-VLAN'",
            "Vérifier l'utilisation du VLAN dans lst_vlan_usage : SELECT * FROM lst_vlan_usage WHERE vlan_id = <vlan_id>",
            "Identifier les équipements encore associés : SELECT eqpt_id FROM t_equipments WHERE vlan_id = <vlan_id>",
            "Vérifier les rôles de production liés : SELECT * FROM t_res_prod_roles WHERE vlan_id = <vlan_id>",
            "Vérifier les liens logiques encore actifs sur ce VLAN",
        ],
        "root_causes": [
            "VLAN encore associé à des équipements actifs dans t_equipments",
            "Ressources de production liées au VLAN non supprimées (t_res_prod_roles)",
            "Utilisation détectée dans lst_vlan_usage — associations résiduelles",
            "Liens logiques (VP/VC) encore actifs sur ce VLAN",
        ],
        "resolution_steps": [
            "1. Récupérer l'identifiant VLAN (vlan_id) dans t_vlans",
            "2. Consulter lst_vlan_usage pour identifier toutes les associations actives",
            "3. Supprimer ou migrer les ressources liées (équipements, rôles de prod)",
            "4. Vérifier et supprimer les liens logiques actifs sur ce VLAN",
            "5. Relancer la suppression du VLAN via IHM BRASIL",
            "6. Valider l'absence d'entrées résiduelles dans lst_vlan_usage",
        ],
        "sql_queries": [
            "SELECT vlan_id, vlan_name, vlan_status FROM t_vlans WHERE vlan_name = 'NOM-VLAN';",
            "SELECT * FROM lst_vlan_usage WHERE vlan_id = <vlan_id>;",
            "SELECT eqpt_id, eqpt_name FROM t_equipments WHERE vlan_id = <vlan_id>;",
            "SELECT * FROM t_res_prod_roles WHERE vlan_id = <vlan_id>;",
        ],
        "tables_involved": [
            "t_vlans", "lst_vlan_usage", "t_equipments",
            "t_res_prod_roles", "t_links", "t_res_prod_controlers",
        ],
        "exceptions_referenced": [],
        "escalation_path": "N3 BRASIL → Equipe BRASIL",
        "trust_score": 0.93,
        "embed_text": (
            "FR 190 suppression VLAN impossible Brasil. "
            "Le VLAN ne peut pas être supprimé via IHM BRASIL. "
            "lst_vlan_usage retourne des associations actives. "
            "t_vlans t_equipments t_res_prod_roles. "
            "Identifier et supprimer les ressources liées au VLAN puis relancer suppression."
        ),
    },

    # ── FR 188 : Suppression BAS/ROUTEUR IMPOSSIBLE ───────────────────────────
    {
        "fr_number": "FR 188",
        "fr_aliases": ["FR188", "FR-188", "FR 188"],
        "glob_pattern": "*188*",
        "proc_id": "FR-188-SUPPRESSION-BAS-ROUTEUR",
        "title": "FR 188 - Suppression d'un BAS ou ROUTEUR impossible",
        "title_normalized": "suppression bas routeur impossible",
        "incident_type": "network_equipment.bas_router_deletion_impossible",
        "application": "BRASIL",
        "applications_involved": ["BRASIL"],
        "estimated_duration": "20 à 45 minutes",
        "sensitivity": "HAUTE — équipements de cœur réseau",
        "symptoms": [
            "Suppression d'un BAS (Broadband Access Server) impossible via IHM BRASIL",
            "Suppression d'un ROUTEUR impossible via IHM BRASIL",
            "Message d'erreur lors de la tentative de suppression",
            "Équipement encore référencé dans des ressources logiques ou des associations",
        ],
        "diagnostic_steps": [
            "Identifier l'équipement dans t_equipments : SELECT eqpt_id, eqpt_name, eqpt_type FROM t_equipments WHERE eqpt_name = 'NOM-BAS'",
            "Vérifier les associations résiduelles dans les tables liées à l'eqpt_id",
            "Vérifier les ressources de production associées dans t_res_prod_controlers",
            "Vérifier les cartes et interfaces actives sur le BAS/routeur",
            "Vérifier les liens logiques encore rattachés à cet équipement",
        ],
        "root_causes": [
            "Données parasites dans t_res_prod_controlers liées à cet équipement",
            "Interfaces ou cartes encore actives non supprimées au préalable",
            "Ressources logiques (VLANs, VPs) encore associées à cet équipement",
            "Liens logiques actifs rattachés au BAS/routeur",
        ],
        "resolution_steps": [
            "1. Vérifier le statut de l'équipement dans t_equipments (eqpt_status='D')",
            "2. Lister toutes les tables référençant l'eqpt_id (requête information_schema)",
            "3. Identifier les données parasites bloquantes",
            "4. Supprimer les données parasites en BDD",
            "5. Relancer la suppression via IHM BRASIL",
            "6. Conserver l'historique des requêtes SELECT/DELETE",
        ],
        "sql_queries": [
            "SELECT eqpt_id, eqpt_name, eqpt_type, eqpt_status FROM t_equipments WHERE eqpt_name = 'NOM-BAS';",
            "SELECT table_name, column_name FROM information_schema.columns "
            "WHERE table_catalog='brasil' AND table_schema='public' AND column_name LIKE '%eqpt_id%' ORDER BY table_name;",
            "SELECT * FROM t_res_prod_controlers WHERE eqpt_id = <eqpt_id>;",
            "DELETE FROM t_res_prod_controlers WHERE eqpt_id = <eqpt_id>;",
        ],
        "tables_involved": [
            "t_equipments", "t_cards", "t_res_prod_controlers",
            "t_links", "t_vlans", "information_schema.columns",
        ],
        "exceptions_referenced": ["BrasilInternalException"],
        "escalation_path": "N3 BRASIL → Equipe BRASIL",
        "trust_score": 0.92,
        "embed_text": (
            "FR 188 suppression BAS routeur impossible Brasil. "
            "Le BAS ou le routeur ne peut pas être supprimé via IHM BRASIL. "
            "Données parasites t_res_prod_controlers bloquantes. "
            "t_equipments t_cards t_links. "
            "Identifier et supprimer les données parasites puis relancer suppression."
        ),
    },

    # ── FR 191 : Suppression en masse de cartes ───────────────────────────────
    {
        "fr_number": "FR 191",
        "fr_aliases": ["FR191", "FR-191", "FR 191"],
        "glob_pattern": "*191*masse*",
        "proc_id": "FR-191-SUPPRESSION-MASSE-CARTES",
        "title": "FR 191 - Suppression en masse de cartes",
        "title_normalized": "suppression masse cartes",
        "incident_type": "network_equipment.bulk_card_deletion",
        "application": "BRASIL",
        "applications_involved": ["BRASIL"],
        "estimated_duration": "30 minutes à 2 heures selon volume",
        "sensitivity": "HAUTE — opération de masse irréversible",
        "symptoms": [
            "Besoin de supprimer un grand nombre de cartes en une seule opération",
            "Suppression unitaire trop longue ou impraticable pour le volume concerné",
            "Cartes à supprimer associées à un DSLAM ou un équipement à retirer du parc",
        ],
        "diagnostic_steps": [
            "Identifier les cartes à supprimer : SELECT card_id, card_name, card_prod_status FROM t_cards WHERE eqpt_id = <eqpt_id>",
            "Vérifier qu'aucune carte n'est encore ouverte à la production (card_prod_status='O')",
            "Lister les dépendances des cartes (ports, liaisons physiques)",
            "Vérifier l'absence de ressources logiques encore associées aux cartes",
        ],
        "root_causes": [
            "Volume important de cartes à supprimer après démantèlement d'équipement",
            "Cartes encore ouvertes à la production bloquant la suppression",
            "Dépendances résiduelles (ports, liaisons) non nettoyées",
        ],
        "resolution_steps": [
            "1. Lister toutes les cartes liées à l'équipement dans t_cards",
            "2. Vérifier card_prod_status='F' (fermé) pour chaque carte",
            "3. Fermer les cartes encore ouvertes via IHM BRASIL si nécessaire",
            "4. Supprimer les cartes en masse via IHM BRASIL (sélection multiple)",
            "5. Vérifier l'absence de données résiduelles après suppression",
        ],
        "sql_queries": [
            "SELECT card_id, card_name, card_prod_status FROM t_cards WHERE eqpt_id = <eqpt_id>;",
            "SELECT card_id FROM t_cards WHERE eqpt_id = <eqpt_id> AND card_prod_status = 'O';",
            "SELECT COUNT(*) FROM t_cards WHERE eqpt_id = <eqpt_id>;",
        ],
        "tables_involved": [
            "t_cards", "t_equipments", "t_shelfs",
            "t_d_dslam_xdsl_cards", "t_ports",
        ],
        "exceptions_referenced": [],
        "escalation_path": "N3 BRASIL → Equipe BRASIL",
        "trust_score": 0.90,
        "embed_text": (
            "FR 191 suppression en masse cartes Brasil. "
            "Supprimer plusieurs cartes simultanément dans BRASIL. "
            "t_cards card_prod_status card_id eqpt_id. "
            "Vérifier cartes fermées avant suppression en masse."
        ),
    },
]


def inject_fr(client: QdrantClient, fr_def: dict, fr_dir: str) -> bool:
    """Indexe une FR dans Qdrant brasil_procedures. Retourne True si succès."""
    fr_number = fr_def["fr_number"]
    proc_id = fr_def["proc_id"]
    point_uuid = str(uuid.uuid5(uuid.NAMESPACE_DNS, proc_id))

    # Vérifier si déjà indexée
    try:
        existing = client.retrieve(
            collection_name="brasil_procedures",
            ids=[point_uuid],
            with_payload=False,
        )
        if existing:
            print(f"[SKIP] {fr_number} déjà indexée (UUID={point_uuid})")
            return True
    except Exception:
        pass

    # Chercher le fichier docx
    pattern = os.path.join(fr_dir, fr_def["glob_pattern"])
    files = glob.glob(pattern, recursive=False)
    if not files:
        # Fallback: pattern plus souple
        fr_num_clean = fr_number.replace(" ", "*")
        files = glob.glob(os.path.join(fr_dir, f"*{fr_num_clean}*"))
    if not files:
        print(f"[WARN] {fr_number} : fichier docx non trouvé (pattern={pattern})")
        full_text = ""
    else:
        docx_path = files[0]
        print(f"[OK] {fr_number} : fichier trouvé: {os.path.basename(docx_path)}")
        full_text, _, _ = extract_docx(docx_path)
        print(f"     Contenu: {len(full_text)} caractères")

    # Construire le payload
    payload = {
        "procedure_id": proc_id,
        "fr_number": fr_number,
        "fr_aliases": fr_def["fr_aliases"],
        "title": fr_def["title"],
        "title_normalized": fr_def["title_normalized"],
        "incident_type": fr_def["incident_type"],
        "application": fr_def.get("application", "BRASIL"),
        "applications_involved": fr_def.get("applications_involved", ["BRASIL"]),
        "estimated_duration": fr_def.get("estimated_duration", ""),
        "sensitivity": fr_def.get("sensitivity", ""),
        "symptoms": fr_def.get("symptoms", []),
        "diagnostic_steps": fr_def.get("diagnostic_steps", []),
        "root_causes": fr_def.get("root_causes", []),
        "resolution_steps": fr_def.get("resolution_steps", []),
        "sql_queries": fr_def.get("sql_queries", []),
        "tables_involved": fr_def.get("tables_involved", []),
        "exceptions_referenced": fr_def.get("exceptions_referenced", []),
        "escalation_path": fr_def.get("escalation_path", "N3 BRASIL"),
        "special_cases": fr_def.get("special_cases", []),
        "source_types": ["fr_document", "catalog_validated"],
        "trust_score": fr_def.get("trust_score", 0.90),
        "ticket_count": 0,
        "cluster_id": f"FR-{fr_number.replace('FR ', '').strip()}",
        "full_text_excerpt": full_text[:2000] if full_text else "",
        "indexed_at": datetime.utcnow().isoformat(),
        "last_updated": datetime.utcnow().isoformat(),
    }

    # Calculer l'embedding
    embedding = compute_embedding(fr_def["embed_text"])

    # Upsert dans Qdrant
    point = PointStruct(id=point_uuid, vector=embedding, payload=payload)
    try:
        result = client.upsert(collection_name="brasil_procedures", points=[point])
        print(f"[OK] {fr_number} indexée — UUID={point_uuid}, status={result.status}")
        return True
    except Exception as e:
        print(f"[ERROR] {fr_number} : upsert échoué: {e}")
        return False


def main():
    print("=" * 60)
    print("inject_fr_batch.py — Indexation batch des FRs dans Qdrant")
    print("=" * 60)

    client = QdrantClient(host="localhost", port=6333)

    # Vérifier que la collection existe
    try:
        info = client.get_collection("brasil_procedures")
        print(f"[OK] Collection brasil_procedures: {info.points_count} points existants")
    except Exception as e:
        print(f"[ERROR] Collection brasil_procedures inaccessible: {e}")
        sys.exit(1)

    # Charger le modèle une fois
    get_embedding_model()

    success_count = 0
    for fr_def in FR_DEFINITIONS:
        print(f"\n--- {fr_def['fr_number']} ---")
        if inject_fr(client, fr_def, FR_DIR):
            success_count += 1

    print(f"\n{'='*60}")
    print(f"Résultat: {success_count}/{len(FR_DEFINITIONS)} FRs indexées avec succès")

    # Vérification finale
    info = client.get_collection("brasil_procedures")
    print(f"Collection brasil_procedures: {info.points_count} points au total")
    print("=" * 60)


if __name__ == "__main__":
    main()
